"""Builds the formatted daily P&L + ad-performance workbook with openpyxl.

Three sheets: a Summary (P&L + blended ad metrics, the one-glance view), an
Orders detail sheet, and an Ad Performance detail sheet — each sourced
straight from the Shopify/Meta client models, no intermediate CSV/manual step.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ecommerce_mcp.clients.meta_ads_client import AdInsight, AdPerformanceSummary
from ecommerce_mcp.clients.shopify_client import DailyPnL, Order

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True)
TOTAL_FONT = Font(bold=True)
CURRENCY_FMT = '"$"#,##0.00'
THIN_BORDER = Border(bottom=Side(style="thin", color="D1D5DB"))


def _style_header_row(ws: Worksheet, row: int, num_cols: int) -> None:
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left")


def _autofit_columns(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_summary_sheet(
    wb: Workbook,
    report_date: date,
    pnl: DailyPnL,
    ad_summary: AdPerformanceSummary,
) -> None:
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = f"Daily P&L + Ad Performance — {report_date.isoformat()}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    rows: list[tuple[str, object, str | None]] = [
        ("", None, None),
        ("Revenue (Shopify)", pnl.revenue, CURRENCY_FMT),
        ("COGS", pnl.cogs, CURRENCY_FMT),
        ("Gross Profit", pnl.gross_profit, CURRENCY_FMT),
        ("Ad Spend (Meta)", ad_summary.spend, CURRENCY_FMT),
        ("Net Profit (Gross Profit − Ad Spend)", pnl.gross_profit - ad_summary.spend, CURRENCY_FMT),
        ("", None, None),
        ("Orders", pnl.order_count, "#,##0"),
        ("Ad Impressions", ad_summary.impressions, "#,##0"),
        ("Ad Clicks", ad_summary.clicks, "#,##0"),
        ("Purchases Attributed (Meta)", ad_summary.purchases, "#,##0"),
        ("Purchase Value (Meta)", ad_summary.purchase_value, CURRENCY_FMT),
        ("Blended ROAS", ad_summary.roas, '0.00"x"'),
    ]

    row_idx = 3
    for label, value, number_format in rows:
        if label == "":
            row_idx += 1
            continue
        label_cell = ws.cell(row=row_idx, column=1, value=label)
        value_cell = ws.cell(row=row_idx, column=2, value=float(value) if isinstance(value, Decimal) else value)
        if number_format:
            value_cell.number_format = number_format
        if label in ("Gross Profit", "Net Profit (Gross Profit − Ad Spend)"):
            label_cell.font = TOTAL_FONT
            value_cell.font = TOTAL_FONT
        row_idx += 1

    _autofit_columns(ws, [36, 16])


def _write_orders_sheet(wb: Workbook, orders: list[Order]) -> None:
    ws = wb.create_sheet("Orders")
    headers = ["Order", "Created At", "Status", "Subtotal", "Discounts", "Tax", "Total"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    for order in sorted(orders, key=lambda o: o.created_at):
        ws.append(
            [
                order.name,
                order.created_at,
                order.financial_status,
                float(order.subtotal_price),
                float(order.total_discounts),
                float(order.total_tax),
                float(order.total_price),
            ]
        )

    for row in ws.iter_rows(min_row=2, min_col=4, max_col=7):
        for cell in row:
            cell.number_format = CURRENCY_FMT

    ws.freeze_panes = "A2"
    _autofit_columns(ws, [10, 24, 12, 12, 12, 10, 12])


def _write_ad_performance_sheet(wb: Workbook, insights: list[AdInsight]) -> None:
    ws = wb.create_sheet("Ad Performance")
    headers = ["Date", "Campaign", "Spend", "Impressions", "Clicks", "Purchases", "Purchase Value", "ROAS"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    for insight in sorted(insights, key=lambda i: (i.date_start, i.campaign_name)):
        ws.append(
            [
                insight.date_start,
                insight.campaign_name,
                float(insight.spend),
                insight.impressions,
                insight.clicks,
                insight.purchases,
                float(insight.purchase_value),
                float(insight.roas) if insight.roas is not None else None,
            ]
        )

    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = CURRENCY_FMT
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        row[0].number_format = CURRENCY_FMT
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].number_format = '0.00"x"'

    ws.freeze_panes = "A2"
    _autofit_columns(ws, [12, 28, 12, 13, 10, 11, 15, 10])


def build_workbook(
    *,
    report_date: date,
    pnl: DailyPnL,
    orders: list[Order],
    ad_summary: AdPerformanceSummary,
    insights: list[AdInsight],
) -> Workbook:
    wb = Workbook()
    _write_summary_sheet(wb, report_date, pnl, ad_summary)
    _write_orders_sheet(wb, orders)
    _write_ad_performance_sheet(wb, insights)
    return wb
